# -*- coding: utf-8 -*-
"""
name_leak_check.py — 원작·레퍼런스·역대본의 인물 이름이 우리 산출물에 남아 있는지 기계로 잡는다.

사용:
  python tools/name_leak_check.py --ref <원작 파일|폴더>... --ours <우리 파일|폴더>... [--extra <이름 목록 txt>] [--allow 이름,이름] [--min 3]

  --ref    원작 텍스트(.txt/.md/.docx). 폴더면 안의 txt/md/docx 전부.
  --ours   우리 산출물(.md/.txt/.docx). 폴더면 안의 md/txt/docx 전부(reference/·_X_·bak 폴더 제외).
  --extra  자동 추출이 못 잡는 이름을 한 줄에 하나씩 적은 파일(한글 음차·별칭).
  --allow  우리 것이 맞는 이름(쉼표 구분) — 원작에도 우연히 같은 흔한 단어가 있을 때만.
  --min    영문 고유명사 최소 등장 횟수(기본 3).

원작 이름 추출 규칙:
  EN  = 화자 줄(줄 전체가 대문자 시작 단어 1~3개) + 본문 대문자 시작 단어 중 소문자로는 한 번도 안 나오는 것(빈도 ≥ --min).
        TITLE CARD "이름 -- 직함" 줄의 이름도 포함. 성·이름을 따로도 검사한다(Reed처럼 성만 남는 사고).
  CN  = "가운뎃점 이름"(杰克·里德)의 조각 + 화자 줄(줄 머리 2~4자 뒤에 ：/:/（).
  KR  = 화자 줄(줄 머리 2~5자 뒤에 :/：/(/（) + 「이름」 형태.
판정 = 우리 파일에 한 건이라도 있으면 exit 1. 문맥을 같이 찍으니 흔한 단어 오탐은 눈으로 걸러 --allow 로 뺀다.
"""
import argparse, os, re, sys, glob, collections

STOP_EN = set("""
INT EXT TITLE CARD EPISODE END CUT FADE SCENE INSERT FLASHBACK MONTAGE VOICE OVER VO OS CONTINUED CONT LATER
The A An And Or But If In On At To Of For By With From Into Onto Over Under Up Down Out Off As Is Are Was Were Be Been
He She It They We You I His Her Its Their Our Your My Me Him Them Us This That These Those There Here Then Than
What When Where Who Why How Which Not No Yes Oh Ah Hey Well Okay Ok Wait Look Stop Go Come Get Let Please Sir Ma'am Mr Mrs Ms Dr
Day Night Daytime Nighttime Morning Evening Afternoon Dawn Dusk Interior Exterior Rain Rainy Sunny Later Same Continuous
City Street Alley Room Office Hospital Hotel Lobby Hall Bedroom Kitchen Car Road Park House Mansion Villa Suite Bath Ward ICU Floor
Mom Dad Mother Father Grandpa Grandma Grandfather Grandmother Son Daughter Brother Sister Uncle Aunt Boss Doctor Nurse Officer Police
Chairman CEO President Director Manager Secretary Assistant Guard Driver Butler Maid Host Guest Customer Waiter Staff
Value Span Life Left Days Years Countdown King Group Corp Inc Ltd Co Company Bank Insurance Hospital
Formatted Script Cast Note Notes Part Act Chapter Page Last Man Family Chart Audience Sec Years Heir Executive Founder
Target Genre Keywords Logline Synopsis Reference Free Paid Total Episode Episodes Vertical Drama Proposal Title
""".split())

def read_text(path):
    if path.lower().endswith(".docx"):
        try:
            from docx import Document
        except ImportError:
            sys.exit("python-docx 필요: pip install python-docx")
        d = Document(path)
        out = [p.text for p in d.paragraphs]
        def walk(tables):
            for tb in tables:
                for row in tb.rows:
                    for c in row.cells:
                        out.append(c.text); walk(c.tables)
        walk(d.tables)
        return "\n".join(out)
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()

def collect(paths, exts):
    files = []
    for p in paths:
        if os.path.isdir(p):
            for root, dirs, fs in os.walk(p):
                dirs[:] = [d for d in dirs if not (d.startswith("_X_") or d in ("reference", "bak", "_archive", ".git", "__pycache__") or "원본" in d or "작업파일" in d)]
                for f in fs:
                    if f.lower().endswith(exts) and not f.startswith("~$"):
                        files.append(os.path.join(root, f))
        elif os.path.exists(p):
            files.append(p)
        else:
            print(f"[경고] 없음: {p}")
    return files

def extract_names(text, min_count):
    names = {"EN": collections.Counter(), "CN": collections.Counter(), "KR": collections.Counter()}
    # --- EN 화자 줄 / TITLE CARD
    for line in text.splitlines():
        s = line.strip()
        m = re.match(r"^((?:[A-Z][a-z]+\.?)(?: [A-Z][a-z]+){0,2})$", s)
        if m:
            names["EN"][m.group(1)] += 1
        m = re.match(r"^TITLE CARD:\s*([A-Z][A-Za-z.]+(?: [A-Z][A-Za-z.]+){0,2})\s*[-—–]", s)
        if m:
            names["EN"][m.group(1)] += 5
    # --- EN 본문 고유명사 (소문자형이 한 번도 없는 대문자 단어)
    caps = collections.Counter(re.findall(r"\b[A-Z][a-z]{2,}\b", text))
    lowers = set(re.findall(r"\b[a-z]{3,}\b", text))
    for w, c in caps.items():
        if c >= min_count and w.lower() not in lowers and w not in STOP_EN:
            names["EN"][w] += c
    # 화자 풀네임은 조각으로도
    for full in list(names["EN"]):
        for piece in full.replace(".", "").split():
            if piece not in STOP_EN and len(piece) >= 3:
                names["EN"][piece] += names["EN"][full]
    # --- CN
    for m in re.finditer(r"([一-鿿]{1,4})·([一-鿿]{1,6})(?:·([一-鿿]{1,6}))?", text):
        for piece in m.groups():
            if piece: names["CN"][piece] += 1
        names["CN"][m.group(0)] += 1
    for line in text.splitlines():
        m = re.match(r"^\s*([一-鿿]{2,4})\s*[：:（(]", line)
        if m: names["CN"][m.group(1)] += 1
    # --- KR
    for line in text.splitlines():
        m = re.match(r"^\s*([가-힣]{2,5})\s*[:：(（]", line)
        if m: names["KR"][m.group(1)] += 1
    for m in re.finditer(r"「([가-힣]{2,5})」", text):
        names["KR"][m.group(1)] += 1
    # 빈도 필터
    out = {}
    out["EN"] = {n for n, c in names["EN"].items() if c >= min_count}
    out["CN"] = {n for n, c in names["CN"].items() if c >= 2 and n not in ("先生", "女士", "小姐", "老板", "妈妈", "爸爸", "医生", "护士", "秘书", "众人", "旁白", "字幕", "画外音", "人物", "主理人", "房主", "字幕卡", "内室", "画廊")}
    out["KR"] = {n for n, c in names["KR"].items() if c >= 3 and n not in ("자막", "내레이션", "장면", "화면", "사람들", "직원", "의사", "간호사", "비서", "경비", "엄마", "아빠", "어머니", "아버지", "할아버지", "할머니", "회장", "사장", "손님", "모두", "일동", "남자", "여자", "경비원", "사모님", "변호사", "남편", "아내", "선장", "선원", "집사", "하녀", "경찰", "기자", "군중", "하객", "종업원", "점원", "웨이터", "운전사", "노인", "소년", "소녀", "아이", "학생", "선생", "교수", "장군", "병사", "시녀", "시종", "무녀", "사제", "왕", "여왕", "왕자", "공주", "황제", "황후")}
    return out

def find_hits(text, names, allow):
    hits = []
    for n in sorted(names["EN"]):
        if n in allow: continue
        for m in re.finditer(r"(?<![A-Za-z])" + re.escape(n) + r"(?![A-Za-z])", text):
            hits.append((n, text[max(0, m.start()-25):m.end()+15].replace("\n", " ")))
    for n in sorted(names["CN"]):
        if n in allow or len(n) < 2 and "·" not in n: continue
        for m in re.finditer(re.escape(n), text):
            hits.append((n, text[max(0, m.start()-15):m.end()+10].replace("\n", " ")))
    for n in sorted(names["KR"]):
        if n in allow: continue
        for m in re.finditer(r"(?<![가-힣])" + re.escape(n) + r"(?=[^가-힣]|[이가은는을를의과와에도만께야아랑로]|$)", text):
            hits.append((n, text[max(0, m.start()-15):m.end()+10].replace("\n", " ")))
    return hits

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ref", nargs="+", required=True)
    ap.add_argument("--ours", nargs="+", required=True)
    ap.add_argument("--extra")
    ap.add_argument("--allow", default="")
    ap.add_argument("--min", type=int, default=3)
    ap.add_argument("--quiet", action="store_true", help="이름 목록 안 찍음")
    a = ap.parse_args()
    allow = {x.strip() for x in a.allow.split(",") if x.strip()}

    refs = collect(a.ref, (".txt", ".md", ".docx"))
    ours = collect(a.ours, (".txt", ".md", ".docx"))
    if not refs: sys.exit("원작 파일 0")
    if not ours: sys.exit("우리 파일 0")

    names = {"EN": set(), "CN": set(), "KR": set()}
    for r in refs:
        ex = extract_names(read_text(r), a.min)
        for k in names: names[k] |= ex[k]
    if a.extra:
        for line in open(a.extra, encoding="utf-8"):
            w = line.strip()
            if not w or w.startswith("#"): continue
            k = "EN" if re.match(r"[A-Za-z]", w) else ("CN" if re.search(r"[一-鿿]", w) else "KR")
            names[k].add(w)
    if not a.quiet:
        print("원작 이름 (자동 추출):")
        for k in ("EN", "CN", "KR"):
            print(f"  {k} {len(names[k])}: " + ", ".join(sorted(names[k])))
        print()

    total = 0
    for o in ours:
        hits = find_hits(read_text(o), names, allow)
        if not hits: continue
        total += len(hits)
        by = collections.defaultdict(list)
        for n, ctx in hits: by[n].append(ctx)
        print(f"🚨 {o}")
        for n, ctxs in sorted(by.items(), key=lambda x: -len(x[1])):
            print(f"   {n} ×{len(ctxs)} | " + " ‖ ".join(ctxs[:2]))
    print()
    if total:
        print(f"FAIL — 원작 이름 잔존 {total}건. 발송·등재 금지. (흔한 단어 오탐이면 --allow 로 빼고 다시)")
        sys.exit(1)
    print(f"PASS — 우리 파일 {len(ours)}개에 원작 이름 0건 (원작 {len(refs)}개 · 이름 EN {len(names['EN'])}/CN {len(names['CN'])}/KR {len(names['KR'])}).")

if __name__ == "__main__":
    main()
