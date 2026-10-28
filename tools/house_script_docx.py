# -*- coding: utf-8 -*-
"""번역 작업공간 하우스 양식 대본 docx — 중문/한국어 2단.

출처 = C:\\Users\\Rowan\\Documents\\script-translator-share-0828\\CLAUDE.md 「번역 출력 형식 규칙」:
- 원문(중문) 줄 바로 다음 줄에 한국어 줄
- 화 제목 · 씬 헤더 = 굵게 · 검정 / 지문 · 인물 줄 · 설정 글 = 보통 · 검정
- 대사 줄(`이름: 대사`) · 표시 줄(【회상】·CUT TO·(E) 등) = 굵게 · 파랑 #1155CC
- 지문↔지문(표시 줄·인물 줄 포함)은 빈 줄 없이 이어 쓰고, 그 밖의 전환은 빈 줄 1
- 화 제목 앞 페이지 나누기(첫 화 제외) · 들여쓰기 0

사용:
    from house_script_docx import build
    items = [('ep', '第1集', '1화'), ('scene', 'S#001-1. 王城外荒野 / 夜 / 外', 'S#001-1. 왕성 밖 황야 / 밤 / 실외'),
             ('cast', '人物：莎拉…', '인물: 사라…'), ('action', '△…', '…'), ('dlg', '莎拉（惊愕）：…', '사라: (놀라서) …')]
    build(items, out_path, comments={item_index: ['말풍선 첫 문단', '둘째 문단']}, title='제목 줄')
kind = 'title' | 'head'(설정 소제목) | 'text'(설정 글) | 'ep' | 'scene' | 'cast' | 'action' | 'mark' | 'dlg'
cn 또는 kr이 빈 문자열이면 그 줄은 넣지 않는다. 말풍선은 한국어 줄(없으면 중문 줄)에 단다.
"""
import docx
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLUE = RGBColor(0x11, 0x55, 0xCC)
BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = '맑은 고딕'
STYLE = {  # kind: (bold, color)
    'title': (True, BLACK), 'head': (True, BLACK), 'text': (False, BLACK),
    'ep': (True, BLACK), 'scene': (True, BLACK), 'cast': (False, BLACK),
    'action': (False, BLACK), 'mark': (True, BLUE), 'dlg': (True, BLUE),
}
FLOW = {'action', 'cast', 'mark', 'text'}   # 이 종류끼리는 빈 줄 없이 이어 쓴다


def _para(d, text, bold, color, size=10.5):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.left_indent = pf.first_line_indent = Pt(0)
    pf.line_spacing = 1.15
    r = p.add_run(text)
    r.bold, r.font.color.rgb, r.font.size, r.font.name = bold, color, Pt(size), FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    return p, r


def build(items, out, comments=None, title=None, author='검토', initials='검'):
    d = docx.Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(1.8)
    if title:
        _para(d, title, True, BLACK, 13)
        _para(d, '', False, BLACK)
    prev, seen_ep, anchor = None, False, {}
    for i, (kind, cn, kr) in enumerate(items):
        bold, color = STYLE[kind]
        if prev is not None and not (prev in FLOW and kind in FLOW):
            _para(d, '', False, BLACK)
        first = True
        for text in (cn, kr):
            if not text:
                continue
            p, r = _para(d, text, bold, color)
            if kind == 'ep' and first and seen_ep:
                p.paragraph_format.page_break_before = True
            first = False
            anchor[i] = r                      # 마지막에 넣은 줄 = 한국어 줄(있으면)
        if kind == 'ep':
            seen_ep = True
        prev = kind
    n = 0
    for i, paras in (comments or {}).items():
        cm = d.add_comment(anchor[i], text=paras[0], author=author, initials=initials)
        for t in paras[1:]:
            cm.add_paragraph(t)
        n += 1
    d.save(out)
    return n
