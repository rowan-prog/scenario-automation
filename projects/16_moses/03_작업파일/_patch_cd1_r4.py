# -*- coding: utf-8 -*-
"""CD1 v2 -> v3: 2026-09-23 원본 docx 대조 r4 (회차 6구간 + 개요/인물) 확정 오류 정정."""
import json, sys
import openpyxl

ROOT = r"C:/Users/Rowan/scenario-automation/projects/16_moses/"
SRC = ROOT + "내 남편은 거지 모세_CD1 셀링 포인트_v2.xlsx"
DST = ROOT + "내 남편은 거지 모세_CD1 셀링 포인트_v3.xlsx"
SPEC = ROOT + "03_작업파일/_cd1_spec.json"

# (cell, old, new)
FIX = [
    # EP1 — 회귀 알아채는 계기 = 귓속말 (매는 그 뒤)
    ("H4", "델릴라가 먼저 황금 새장의 매를 잡아 파라오에게 날려 보내는 것을 보며 델릴라도 같은 날로 돌아왔음을 알아챈다. 매가 파라오의 어깨에 앉고 가족들이 환호한다.",
           "이번 생의 파라오는 자기가 갖겠다는 델릴라의 귓속말을 듣고 델릴라도 같은 날로 돌아왔음을 알아챈다. 델릴라가 먼저 황금 새장의 매를 잡아 파라오에게 날려 보내고, 매가 파라오의 어깨에 앉자 가족들이 환호한다."),
    # EP2 — 모래는 노예 구역 가기 전에 사라짐
    ("H5", "비둘기가 황금빛 모래에 싸여 노예 구역으로 날아간다.",
           "비둘기 주위로 황금빛 모래가 반짝이다 사라지고, 비둘기는 노예 구역으로 날아간다."),
    # EP4 — 델릴라 혼사를 다행이라 한 건 네페라
    ("H7", "라반은 다말을 다그치며 델릴라의 행운을 칭찬한다.",
           "라반은 다말을 쓸모없는 것이라 다그치고, 네페라는 델릴라라도 신의 피를 이은 가문과 맺어져 다행이라며 라반을 달랜다."),
    # EP11 — 등장은 10화 끝
    ("H14", "델릴라와 파라오가 함께 다말 앞에 나타나 남편을 조롱하고,",
            "델릴라와 파라오가 혼자 온 다말의 남편을 조롱하고,"),
    # EP15 장면 — 조롱이 먼저
    ("G18", '[장면] 다말이 신발을 벗고 불타는 장작길에 맨발로 들어선다. 델릴라는 불길 위에서 펜던트를 흔든다. "walk the sacred path of fire barefoot."',
            '[장면] 델릴라가 불길 위에서 펜던트를 흔든다. "walk the sacred path of fire barefoot." 다말이 신발을 벗고 불타는 장작길에 맨발로 들어선다.'),
    # EP30 — 병사에게 붙들린 채 몸부림
    ("H33", "다말이 몸부림치며 막아서지만,",
            "병사들에게 붙들린 다말이 울며 몸부림치지만,"),
    # EP31
    ("H34", "이마가 찢어지고,", "이마를 계단 모서리에 부딪혀 피가 흐르고,"),
    # EP41 — 보는 대상 = 백성들
    ("H44", "다말과 모세는 서로에게 기대어 그 광경을 지켜본다.",
            "다말과 모세는 서로에게 기대어 백성들을 바라본다."),
    # EP42 — 고백은 모세만
    ("H45", "사막을 걷던 모세와 다말은 서로에게 다정한 마음을 고백하다가, 모세가 무언가를 알아채고 뒤를 돌아본다.",
            "사막을 걷던 모세가 다말에게 다정한 마음을 고백하다가 무언가를 알아채고 뒤를 돌아본다."),
    # EP49 — 팔찌 먼저, 절규는 그 뒤 / 대사 화자 명시
    ("G52", '다말이 그 손목에서 대왕비 팔찌를 빼내 웅덩이에 던진다. "Just kill me!"',
            '다말이 그 손목에서 대왕비 팔찌를 빼내 웅덩이에 던진다. 델릴라가 울부짖는다. "Just kill me!"'),
    ("H52", "델릴라가 잔해를 타고 기어와 애원하자 다말은 그녀의 손목에서 대왕비의 팔찌를 빼앗아 웅덩이에 던지고, 죽지 말고 죄값을 치르며 살라고 말한다.",
            "델릴라가 잔해를 타고 다말의 발치까지 기어오자 다말은 그녀의 손목에서 대왕비의 팔찌를 빼앗아 웅덩이에 던지고, 델릴라가 차라리 죽이라고 울부짖자 죽지 말고 살아서 죗값을 치르라고 말한다."),
    ("B45", '다말이 델릴라의 손목에서 대왕비의 팔찌를 빼낸다. "Give it back! That\'s the Great Royal Wife\'s bracelet!" 다말이 팔찌를 웅덩이에 던지고 귓가에 말한다. "Live, and pay the price."',
            '다말이 델릴라의 손목에서 대왕비의 팔찌를 빼낸다. 델릴라가 소리친다. "Give it back! That\'s the Great Royal Wife\'s bracelet!" 다말은 말없이 팔찌를 웅덩이에 던진다. 델릴라가 차라리 죽이라고 울부짖자, 다말이 몸을 일으키며 말한다. "Live, and pay the price."'),
    ("C45", 'Tamar rips the Great Royal Wife\'s bracelet off her wrist. "Give it back! That\'s the Great Royal Wife\'s bracelet!" Tamar throws it into a puddle and bends to her ear: "Live, and pay the price."',
            'Tamar rips the Great Royal Wife\'s bracelet off her wrist, and Delilah screams, "Give it back! That\'s the Great Royal Wife\'s bracelet!" Tamar throws it into a puddle without a word. Delilah shrieks at her to just kill her, and Tamar straightens up: "Live, and pay the price."'),
    ("D45", "塔玛从她手腕上扯下大王妃手镯。“还给我！那是大王妃的手镯！”塔玛把手镯扔进水洼，俯身在她耳边说：“活着，付出代价。”",
            "塔玛从她手腕上扯下大王妃手镯，德莉拉尖叫：“还给我！那是大王妃的手镯！”塔玛一言不发地把手镯扔进水洼。德莉拉嘶喊着让她干脆杀了自己，塔玛直起身说：“活着，付出代价。”"),
    # B17 영문 — 낙인은 찍히기 직전
    ("C17", "the branding iron and then the sandstorm,",
            "the sandstorm right before the branding iron lands,"),
    # 다말 — 굴욕마다 왕궁 재앙 아님
    ("B23", "그리고 그 뒤에 왕궁에 재앙이 떨어진다.",
            "그리고 다말을 괴롭힌 사람들에게 재앙이 하나씩 떨어진다."),
    ("C23", "And after each one, a plague hits the palace.",
            "And one by one, plagues fall on the people who hurt her."),
    ("D23", "而在那之后，灾祸便降临王宫。",
            "而欺负她的人，一个接一个遭了灾祸。"),
    # 델릴라 — 낙인은 닿지 않음 / 재앙은 회차별로
    ("B25", "달군 낙인을 얼굴에 대며(EP15~16)", "달군 낙인을 얼굴에 내리찍으려 하고(EP15~16)"),
    ("B25", "그때마다 개구리와 이가 몸에 들러붙고 우박에 맞는다.",
            "개구리가 방을 뒤덮고(EP24), 발을 씻긴 새벽에는 이가 몸에 들러붙고(EP28), 광장에서는 우박에 맞는다(EP30)."),
    ("C25", "holds a red-hot brand to her face (EP15-16)", "tries to bring a red-hot brand down on her face (EP15-16)"),
    ("C25", "Each time, frogs and lice crawl over her and hail hits her.",
            "Frogs swarm her room (EP24), lice crawl over her the dawn after the foot-washing (EP28), and hail hits her in the square (EP30)."),
    ("D25", "把烧红的烙铁按向她的脸（第15~16集）", "举起烧红的烙铁要往她脸上按（第15~16集）"),
    ("D25", "每一次，青蛙和虱子都爬满她的身体，冰雹砸中她。",
            "青蛙爬满她的房间（第24集），洗脚后的黎明虱子爬满她全身（第28集），在广场上又被冰雹砸中（第30集）。"),
    # 라반 — 침묵 동조 = 13화(불의 길에선 경악) / 감옥 = 26화
    ("B27", "다말의 뺨을 때려 쓰러뜨리고(EP13), 델릴라가 언니를 불 위로 걷게 할 때 침묵으로 편든다. 파라오에게 붙잡혀 지하 감옥에 갇히고 광장에서 무릎 꿇려 맞는다(EP30).",
            "침묵으로 델릴라 편을 들고 다말의 뺨을 때려 쓰러뜨린다(EP13). 파라오에게 붙잡혀 지하 감옥에 갇히고(EP26), 광장에서 무릎 꿇려 맞는다(EP30)."),
    ("C27", "He slaps her to the ground (EP13) and sides with Delilah in silence while she makes her sister walk on fire. Pharaoh throws him in the dungeon and has him beaten on his knees in the square (EP30).",
            "He sides with Delilah in silence and slaps Tamar to the ground (EP13). Pharaoh throws him in the dungeon (EP26) and has him beaten on his knees in the square (EP30)."),
    ("D27", "他一巴掌把塔玛打倒在地（第13集），德莉拉逼姐姐走火路时他以沉默站在小女儿一边。后被法老关进地牢，在广场上跪着挨打（第30集）。",
            "他以沉默站在小女儿一边，又一巴掌把塔玛打倒在地（第13集）。后被法老关进地牢（第26集），在广场上跪着挨打（第30集）。"),
    # 페이크 설명 — 15화 뒤는 재앙 아닌 직접 응징
    ("A69", "15·28화의 굴욕은 그 뒤 다말이 가족과 왕궁에 붙잡혀 당하는 일이며, 그때마다 모세의 재앙이 상대에게 떨어진다.",
            "15·28화의 굴욕은 그 뒤 다말이 가족과 왕궁에 붙잡혀 당하는 일이다. 15화 뒤에는 모세가 직접 와서 델릴라를 걷어차고(17화), 28화 뒤에는 이가 델릴라에게 들러붙는다."),
]

wb = openpyxl.load_workbook(SRC)
ws = wb.worksheets[1]
bad = []
for cell, old, new in FIX:
    v = ws[cell].value or ""
    if v.count(old) != 1:
        bad.append((cell, v.count(old), old[:30]))
        continue
    ws[cell].value = v.replace(old, new)
if bad:
    print("MISMATCH", bad); sys.exit(1)
wb.save(DST)
print("xlsx ok", len(FIX), "edits ->", DST)

# spec 사본 동기화 (문자열 전부에서 치환, 건수만 보고)
spec = json.load(open(SPEC, encoding="utf-8"))
hits = {}
def walk(o):
    if isinstance(o, dict):
        return {k: walk(x) for k, x in o.items()}
    if isinstance(o, list):
        return [walk(x) for x in o]
    if isinstance(o, str):
        for cell, old, new in FIX:
            if old in o:
                hits[(cell, old[:20])] = hits.get((cell, old[:20]), 0) + 1
                o = o.replace(old, new)
    return o
spec = walk(spec)
json.dump(spec, open(SPEC, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
print("spec synced", len(hits), "/", len(FIX))
missing = [c for c, o, n in FIX if (c, o[:20]) not in hits]
print("spec not found:", missing)
