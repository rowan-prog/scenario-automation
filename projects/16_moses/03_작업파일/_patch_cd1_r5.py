# -*- coding: utf-8 -*-
"""CD1 v3 -> v4: 2026-09-23 r5. 사용자 지적(B34 '~자리' 비한국어·빈약 / 전문 부자연 표현 / 서구권 기독교 소구) 반영.
원작 사실 = 01_원본/16_moses_source_apollo.md, 각색 사실 = 03_작업파일/_cd1_script.md 원문 확인분만."""
import sys
import openpyxl

ROOT = r"C:/Users/Rowan/scenario-automation/projects/16_moses/"
SRC = ROOT + "내 남편은 거지 모세_CD1 셀링 포인트_v3.xlsx"
DST = ROOT + "내 남편은 거지 모세_CD1 셀링 포인트_v4.xlsx"

# 칸 전체 교체
SET = {
    "B34": (
        "동생 에리스와 함께 회귀한 아이린이 조용히 살려고 아무도 없는 쪽으로 화관을 던지는데, 에리스가 등을 떠밀어 "
        "화관이 남루한 떠돌이 음유시인의 손에 떨어지는 장면. 지난 생에 에리스가 엮였던 바로 그 남자다. "
        "혼례 날 에리스의 지참금으로 금화 아흔아홉 상자가 공개되고 천 주머니 하나뿐인 아이린이 비웃음을 사는데, "
        "오지 않던 신랑이 \"내 신부를 맞으러 왔다\"는 목소리와 함께 나타나는 장면. "
        "아이린을 괴롭히던 가족들이 그 음유시인이 빛의 신 아폴론이라는 걸 먼저 알고 바닥에 엎드리는 장면. "
        "아폴론의 어머니 레토 여신이 둘의 혼인을 반대하며 아이린을 신전 하녀로 시험하고, 아폴론은 정체도 힘도 쓰지 못한 채 "
        "경비병으로 변장해 곁을 지키는 장면. 모든 시련을 견딘 아이린이 그제야 남편이 아폴론이라는 걸 알게 되는 장면(44화). "
        "끝까지 언니를 해치려던 에리스가 용암 감옥에 사슬로 묶인 채 언니와 마주하는 장면(46화).\n"
        "각색은 이 흐름을 그대로 따르면서 화관을 비둘기로, 아폴론을 예언자 모세로, 신의 힘을 열 가지 재앙과 바다 가르기로 바꿨다. "
        "모세와 출애굽, 열 가지 재앙은 기독교 문화에 익숙한 서구권 시청자라면 설명 없이도 아는 이야기라, 그리스 신화보다 쉽게 다가간다."
    ),
    "B16": (
        "나를 비웃던 가족과 왕 앞에서, 내가 고른 거지가 바다를 가른다. "
        "나를 괴롭힌 동생의 몸에는 개구리와 이가 들러붙고, 머리 위로 우박이 떨어진다. "
        "내 자리를 빼앗은 동생이 진흙투성이로 내 발치까지 기어와 울부짖는다."
    ),
    "C16": (
        "The beggar you chose parts the sea in front of the family and the king who laughed at you. "
        "Frogs and lice crawl over the sister who tormented you, and hail comes down on her head. "
        "The sister who stole your place crawls through the mud to your feet, wailing."
    ),
    "B17": (
        "굴욕 장면 바로 뒤에 그 값을 치르는 장면이 붙어 있다. 발을 씻긴 뒤에는 이 떼가, 라반이 맞은 뒤에는 우박이 오고, "
        "얼굴에 낙인이 닿기 직전에는 모래폭풍이 막아선다. 18화에서는 델릴라와 파라오의 얼굴이 다말이 걸었던 불길 앞으로 끌려간다. "
        "바다는 두 번 갈라진다. 7화 혼례 날 지참금 망신 바로 뒤에 한 번, 46~47화 백성이 바다를 건널 때 또 한 번. "
        "델릴라 몸에 개구리와 이가 들러붙는 장면은 우습게 쓰여 있다. 델릴라가 무너지는 장면은 34·35·49화에 세 번 나온다."
    ),
    "C17": (
        "Every humiliation is followed right away by its payback: lice after the foot-washing, hail after Laban is beaten, "
        "and a sandstorm right before the branding iron lands. In EP18 Delilah and Pharaoh have their faces dragged to the same fire Tamar walked through. "
        "The sea parts twice: once in EP7, right after the dowry humiliation on the wedding day, and again in EP46-47 when the people cross. "
        "The frogs and lice crawling over Delilah are written for laughs. Delilah's fall happens three separate times, in EP34, EP35 and EP49."
    ),
}

# (cell, old, new) 부분 교체
FIX = [
    ("B15", "가족이 비웃는 앞에서 스스로 거지를 남편으로 고른 여자가, 혼례 날 그 거지가 바다를 갈라 자기를 데리러 오는 것을 본다.",
            "온 가족이 비웃는데도 거지를 남편으로 고른 여자에게, 혼례 날 그 거지가 바다를 가르며 데리러 온다."),
    ("B18", "히든 아이덴티티 로맨스를 보는 층.", "히든 아이덴티티 로맨스를 즐겨 보는 시청자."),
    ("B19", "모세와 출애굽 세계로 옮긴 각색. 노출과 정사는 9화 첫날밤 한 번(남주 상의 탈의, 드레스 흘러내림, 키스까지). 그 밖의 키스는 10화(이마)·23·40·50화.",
            "모세와 출애굽 이야기로 바꾼 각색이다. 노출과 정사 장면은 9화 첫날밤 한 번뿐이다(남주가 상의를 벗고 드레스가 흘러내린 뒤, 입 맞추며 침대로 쓰러지고 실루엣으로 넘어간다). 그 밖의 키스는 10화(이마)·23화·40화·50화에 있다."),
    ("C19", "(ML shirtless, dress slipping off, kiss)",
            "(ML shirtless, dress slipping off, they fall onto the bed kissing and it cuts to silhouettes)"),
    # 다말 — 비둘기는 빈 쪽으로 날렸다(EP2) / '자리' / 팔리는 그림
    ("B23", "이번에는 가족이 보는 앞에서 비둘기를 노예 구역으로 날려 거지를 남편으로 고른다. 뺨을 맞고, 맨발로 불 위를 걷고, 동생의 발을 씻기고, 목에 칼이 닿는 자리에서도 남편을 부르지 않는다.",
            "이번에는 아무도 없는 쪽으로 비둘기를 날리는데, 비둘기가 노예 구역의 거지 모세에게 가 앉자 가족 앞에서 그를 남편으로 받아들인다. 뺨을 맞고, 맨발로 불 위를 걷고, 동생의 발을 씻기고, 목에 칼이 들어와도 남편을 부르지 않는다."),
    ("B23", "이 작품에서 팔리는 그림은 전부 이 인물에게서 나온다.", "광고로 쓸 만한 장면은 대부분 다말에게서 나온다."),
    ("C23", "This time she sends her dove to the slaves' quarter in front of her whole family and marries a beggar.",
            "This time she sends her dove toward an empty corner, it lands with the beggar Moses in the slaves' quarter, and she takes him as her husband in front of her family."),
    ("C23", "Every sellable image in this title comes from her.", "Most of the ad-worthy scenes come from her."),
    ("D23", "这一世她当着全家人的面把鸽子放向奴隶区，选了一个乞丐做丈夫。",
            "这一世她把鸽子放向没人的地方，鸽子却落到奴隶区的乞丐摩西身边，她当着全家人的面认他做丈夫。"),
    ("D23", "本作所有能卖的画面，全部来自这个人物。", "能用来做广告的画面，大多来自这个人物。"),
    # 모세 — 바다 갈림 시작 = 46화
    ("B24", "(EP47~49)", "(EP46~49)"),
    ("C24", "(EP47-49)", "(EP46-49)"),
    ("D24", "（第47~49集）", "（第46~49集）"),
    # 델릴라 — 비유 정리
    ("B25", "전생에서 거지와 엮여 밑바닥에서 썩다가 언니를 죽였고,", "전생에 거지와 엮여 비참하게 살다가 언니를 죽였고,"),
    ("B25", "언니를 진흙에 처박는 것이 목표다.", "언니를 자기보다 더 비참하게 만드는 것이 목표다."),
    ("B25", "죽지 못하고 살아서 대가를 치른다.", "죽지도 못한 채 살아서 죗값을 치른다."),
    # 파라오 — 개구리 약속 회차
    ("B26", "개구리가 사라지자마자 말을 바꾸며,", "개구리가 사라지자마자 말을 바꾸며(EP24),"),
    ("C26", "breaks his word the moment the frogs are gone,", "breaks his word the moment the frogs are gone (EP24),"),
    ("D26", "却在青蛙一消失就翻脸，", "却在青蛙一消失就翻脸（第24集），"),
    # 라반
    ("B27", "딸은 가문의 영광을 사 오는 물건이다.", "딸을 가문의 영광을 가져올 도구로만 여긴다."),
    ("B27", "혼인하는 순간 가문과 남남이라고 선언한다(EP5).", "거지와 혼인하면 가문에서 내치겠다고 선언한다(EP5)."),
    ("B27", "다말에게 네 아비가 아니냐고 매달린다(EP39).", "다말에게 그래도 내가 네 아비가 아니냐며 매달린다(EP39)."),
    # 회차 요약 문장
    ("H13", "모세는 다말에게 양치기 시절 동료라 둘러대고 아론과 함께 자리를 만든다.",
            "모세는 다말에게 양치기 시절 동료들이 아론을 찾아온 것이라 둘러대고, 아론도 얼른 말을 맞춘다."),
    ("H14", "델릴라와 파라오가 혼자 온 다말의 남편을 조롱하고,", "델릴라와 파라오는 다말이 혼자 왔다며 남편을 조롱하고,"),
    ("H15", "성소에서 다말 앞으로 보물 상자가 도착하자", "성소 수행원들이 다말에게 보내는 보물 상자들을 저택 앞에 늘어놓자"),
    ("H16", "수풀 뒤에 숨어 지켜보던 아론이 겁에 질린 채 지켜본다.", "수풀 뒤에 숨은 아론이 겁에 질린 채 이를 지켜본다."),
    ("H22", "용기를 얻은 파라오가 병사의 단검을 빼앗아 모세를 겨누며 왕국의 정의를 묻는 모세의 말에 병사들에게 반역자를 모조리 죽이라 명령한다.",
            "용기를 얻은 파라오가 병사의 단검을 빼앗아 든다. 모세가 이것이 네 왕국의 정의냐고 묻자, 파라오는 모세에게 칼을 겨누며 반역자들을 모조리 죽이라고 명령한다."),
    ("H23", "다말이 모세를 위험에서 지키려 스스로 몸을 일으켜 파라오 앞으로 절뚝이며 걸어가 자신을 잡고 남편에게 자비를 베풀어달라 애원하지만, 힘이 다해 쓰러진다.",
            "다말이 모세의 품에서 내려와 파라오 앞으로 절뚝이며 걸어가, 모든 게 자기 탓이니 자신을 잡아가고 남편은 살려 달라고 애원하다 힘이 다해 쓰러진다."),
    ("H24", "델릴라가 검으로 모세의 어깨를 찌르지만, 모세는 펜던트를 다말에게 넘긴 채로도 상처만 입을 뿐 쓰러지지 않는다.",
            "델릴라가 검으로 모세의 어깨를 찌르지만, 모세는 자신을 지켜 주던 펜던트가 다말에게 가 있는데도 검을 뽑아 내던지고 버틴다."),
    ("H50", "모세는 행렬의 맨 뒤를 지키며 모두가 지나갈 때까지 물벽을 지킨다.", "모세는 모두가 지나갈 때까지 행렬 맨 뒤에 남는다."),
]

wb = openpyxl.load_workbook(SRC)
ws = wb.worksheets[1]
bad = []
for cell, old, new in FIX:
    v = ws[cell].value or ""
    if v.count(old) != 1:
        bad.append((cell, v.count(old), old[:30]))
        continue
    ws[cell].value = v.replace(old, new)
for cell, new in SET.items():
    if not ws[cell].value:
        bad.append((cell, "empty"))
    ws[cell].value = new
if bad:
    print("MISMATCH", bad); sys.exit(1)
wb.save(DST)
print("ok", len(FIX), "fixes +", len(SET), "sets ->", DST)
