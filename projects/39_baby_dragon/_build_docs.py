# -*- coding: utf-8 -*-
"""39번 작가 전달 docx 빌더 (44번 빌더 복사).
  py -I -X utf8 _build_docs.py proposal <md> <out.docx>
  py -I -X utf8 _build_docs.py script   <md> <out.docx>
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

FONT = 'Malgun Gothic'
GRAY = RGBColor(0x5a, 0x5a, 0x5a)
LGRAY = RGBColor(0x64, 0x64, 0x64)


def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for m in ('left_margin', 'right_margin'):
        setattr(s, m, Cm(2.2))
    s.top_margin = s.bottom_margin = Cm(2.0)
    st = d.styles['Normal']
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.25
    return d


def add(d, segs, size=10, bold=False, color=None, after=3, before=0, page_break=False):
    """segs = str 또는 [(text, bold)] — **굵게** 마크다운은 호출 전에 풀어 넣는다."""
    p = d.add_paragraph()
    if page_break:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if isinstance(segs, str):
        segs = [(segs, bold)]
    for text, b in segs:
        r = p.add_run(text)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        r.font.size = Pt(size)
        r.font.bold = b
        if color is not None:
            r.font.color.rgb = color
    return p


def md_segs(line):
    out = []
    for i, part in enumerate(re.split(r'\*\*', line)):
        if part:
            out.append((part.replace('`', ''), i % 2 == 1))
    return out


def build_proposal(md, out):
    d = new_doc()
    for ln in open(md, encoding='utf-8').read().split('\n'):
        s = ln.rstrip()
        if not s:
            continue
        if s.startswith('# '):
            add(d, s[2:], size=15, bold=True, after=10)
        elif s.startswith('## '):
            add(d, s[3:], size=12, bold=True, before=10, after=4)
        elif re.fullmatch(r'\*\*\d+화\*\*', s):
            add(d, s.strip('*'), size=10.5, bold=True, before=6, after=2)
        elif s.startswith('- '):
            add(d, md_segs('· ' + s[2:]))
        elif s.startswith('  '):
            add(d, md_segs(s.strip()), color=GRAY, size=9.5)
        else:
            add(d, md_segs(s))
    d.save(out)


def build_script(md, out):
    d = new_doc()
    lines = open(md, encoding='utf-8').read().split('\n')
    cur_ep = None
    for ln in lines:
        s = ln.rstrip()
        if not s:
            continue
        if s.startswith('# '):
            add(d, s[2:], size=15, bold=True, after=6)
            continue
        m = re.match(r'## EP(\d+) - (S#\d+.*)', s)
        if m:
            ep = int(m.group(1))
            if ep != cur_ep:
                add(d, f'{ep}화', size=13, bold=True, page_break=cur_ep is not None, after=6)
                cur_ep = ep
            add(d, m.group(2), size=11, bold=True, before=8, after=2)
            continue
        if s.startswith('등장인물:'):
            add(d, s, size=9, color=LGRAY, after=4)
        elif s.startswith('△'):
            add(d, s, color=GRAY)
        elif s.startswith('자막:'):
            add(d, s, color=GRAY)
        elif s.startswith('>'):
            add(d, s.lstrip('> '), size=9.5, color=LGRAY, after=8)
        else:
            add(d, s)
    d.save(out)


if __name__ == '__main__':
    kind, md, out = sys.argv[1:4]
    {'proposal': build_proposal, 'script': build_script}[kind](md, out)
    print('saved', out)
