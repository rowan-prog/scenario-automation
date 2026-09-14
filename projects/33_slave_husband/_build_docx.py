# -*- coding: utf-8 -*-
"""33_slave_husband_p0_rough.md -> docx (작가 발송용 · 페이지 그대로)"""
import re, sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

SRC = "33_slave_husband_p0_rough.md"
OUT = "밤마다 나를 탐하는 노예 남편_무료회차 트리트먼트_워싱 v1.docx"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "맑은 고딕"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

def add(text, bold=False, size=None, space=4):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space)
    return p

for raw in open(SRC, encoding="utf-8").read().splitlines():
    line = raw.rstrip()
    if not line.strip():
        continue
    if line.startswith("### "):
        add(line[4:].strip(), bold=True, size=16, space=10); continue
    if re.match(r"^EP\d+$", line.strip()) or re.match(r"^\d+화$", line.strip()):
        add(line.strip(), bold=True, size=11, space=2); continue
    if line.startswith("[VER"):
        add(line.strip(), bold=True, size=12, space=6); continue
    m = re.match(r"^\*\*(.+?)\*\*\s*(.*)$", line)
    if m:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        r = p.add_run(m.group(1)); r.bold = True
        if m.group(2): p.add_run("  " + m.group(2))
        continue
    if line.startswith("- ") or re.match(r"^\d+\.\s", line):
        add(line.strip()); continue
    add(line.strip())

doc.save(OUT); print("saved", OUT)
