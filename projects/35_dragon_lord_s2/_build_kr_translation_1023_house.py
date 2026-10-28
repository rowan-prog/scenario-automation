# -*- coding: utf-8 -*-
"""1023판 1~8화 한국어 번역본 -> 하우스 양식 docx. 번역 문장은 kr_body.txt 그대로, 양식만 바꾼다."""
import os, re, sys, ast
import docx
sys.path.insert(0, r'C:\Users\Rowan\scenario-automation\tools')
from house_script_docx import build

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = r'C:\Users\Rowan\Downloads\【1-8】《ONE NIGHT WITH THE DRAGON LORD：THE DRAGON EGGS》1023.docx'
OUT = os.path.join(HERE, 'One Night with the Dragon Lord - The Dragon Eggs_작가 대본(1023 · 1-8화)_한국어 번역본(하우스 양식).docx')
TITLE = 'One Night with the Dragon Lord: The Dragon Eggs — 작가 대본 1023 · 1~8화 · 한국어 번역본'
cn = [p.text.strip() for p in docx.Document(SRC).paragraphs]
kr = open(os.path.join(HERE, '35_dragon_lord_s2_writer_1023_kr_body.txt'), encoding='utf-8').read().split('\n')[:328]
assert len(cn) == len(kr) == 328
# COMMENTS 리터럴을 기존 빌더에서 그대로 가져온다
src = open(os.path.join(HERE, '_build_kr_translation_1023.py'), encoding='utf-8').read()
m = re.search(r'^COMMENTS = (\[.*?^\])', src, re.S | re.M)
COMMENTS = ast.literal_eval(m.group(1))
assert len(COMMENTS) == 42, len(COMMENTS)

ERR_DLG = {41, 72}      # 원고 쌍점 오류 줄: 한국어 그대로
ACTION_FORCE = {269}    # 대사가 지문 줄 안에 붙은 곳
DLG_CN = re.compile(r'^[^：△【（(]{1,14}(（[^）]*）)?(OS|VO)?：')
DLG_KR = re.compile(r'^([^:(]{1,20}?)(?:\(([^)]*)\))?: (.*)$')

def kr_dlg(s):
    m = DLG_KR.match(s)
    assert m, s
    name, d, text = m.groups()
    tag = ''
    if d:
        mm = re.search(r'\s*\b(OS|VO)\b\s*$', d)
        if mm:
            tag = '(%s)' % mm.group(1); d = d[:mm.start()].strip()
            if d == '' or d in ('OS', 'VO'): d = ''
        elif d in ('OS', 'VO'):
            tag = '(%s)' % d; d = ''
    return '%s%s: %s%s' % (name, tag, '(%s) ' % d if d else '', text)

def scene(c, k):
    m = re.match(r'^(\d-\d)\s*(.*?)\s+(日|夜)\s+(内|外)$', c)
    n, place, t, io = m.groups()
    sp = '' if c[len(n)] != ' ' else ' '
    cn_l = 'S#%s.%s%s / %s / %s' % (n, ' ' if sp else '', place, t, io)
    m2 = re.match(r'^(\d-\d)\s*(.*?)\s*/\s*(낮|밤)\s*/\s*(내부|외부)$', k)
    kn, kp, kt, kio = m2.groups()
    kr_l = 'S#%s. %s / %s / %s' % (kn, kp, kt, '실내' if kio == '내부' else '실외')
    return cn_l, kr_l

items, idx = [], {}
for i, (c, k) in enumerate(zip(cn, kr)):
    if not c:
        assert not k, i
        continue
    if i == 0: it = ('head', c, k)
    elif re.match(r'^第\s*\d+\s*集$', c):
        n = int(re.search(r'\d+', c).group())
        it = ('ep', c + ('（付费卡点）' if n == 8 else ''), '%d화%s' % (n, ' (페이월)' if n == 8 else ''))
    elif re.match(r'^\d-\d', c): it = ('scene',) + scene(c, k)
    elif c.startswith('出场人物'): it = ('cast', c, re.sub(r'^등장인물: ', '인물: ', k))
    elif re.match(r'^(△)?【(闪回|回忆)', c) and c.endswith('】') or c.startswith('【字幕'):
        it = ('mark', c, k.lstrip('△'))
    elif i in ERR_DLG: it = ('dlg', c, k)
    elif i in ACTION_FORCE: it = ('action', c, k.lstrip('△'))
    elif not c.startswith('△') and DLG_CN.match(c): it = ('dlg', c, kr_dlg(k))
    else: it = ('action', c, k.lstrip('△'))
    idx[i] = len(items); items.append(it)

comments, extra, seen = {}, [], set()
for p, paras in COMMENTS:       # build()는 항목당 말풍선 1개 -> 같은 줄의 두 번째 말풍선은 저장 후 덧붙인다
    if idx[p] in seen: extra.append((p, paras))
    else: seen.add(idx[p]); comments[idx[p]] = paras
n = build(items, OUT, comments=comments, title=TITLE)
if extra:
    d = docx.Document(OUT)
    for p, paras in extra:
        kr_text = items[idx[p]][2]
        para = [x for x in d.paragraphs if x.text == kr_text][0]
        cm = d.add_comment(para.runs[-1], text=paras[0], author='검토', initials='검')
        for t in paras[1:]: cm.add_paragraph(t)
        n += 1
    d.save(OUT)
print('items', len(items), 'comments', n, 'out', OUT)

if __name__ == '__main__' and '--check' in sys.argv:
    from collections import Counter
    print([p for p, c in Counter(p for p, _ in COMMENTS).items() if c > 1])
    for k, c, t in items:
        if k == 'dlg' and not re.match(r'^[^:]{1,20}: ', t): print('DLGEXC', t)
        if k == 'dlg' and ('OS' in t[:20] or 'VO' in t[:20] or '(' in t[:25]): print(t[:60])
    print(Counter(k for k, _, _ in items))
