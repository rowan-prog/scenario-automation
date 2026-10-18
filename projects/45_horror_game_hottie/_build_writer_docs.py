# -*- coding: utf-8 -*-
"""45번 중국 작가 전달 docx 2종 빌더 (중문).
  py -I -X utf8 _build_writer_docs.py proposal <md> <out.docx> <이미지 폴더>
  py -I -X utf8 _build_writer_docs.py script   <md> <out.docx>
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = 'Microsoft YaHei'
GRAY = RGBColor(0x5a, 0x5a, 0x5a)
LGRAY = RGBColor(0x70, 0x70, 0x70)
RED = RGBColor(0xC0, 0x00, 0x00)


def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2.0)
    st = d.styles['Normal']
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.3
    return d


def shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def add(d, segs, size=10, bold=False, color=None, after=3, before=0,
        page_break=False, align=None, fill=None, keep_next=False):
    p = d.add_paragraph()
    if page_break:
        p.paragraph_format.page_break_before = True
    if keep_next:
        p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align:
        p.alignment = align
    if fill:
        shade(p, fill)
    if isinstance(segs, str):
        segs = [(segs, bold)]
    for text, b in segs:
        r = p.add_run(text)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        r.font.size = Pt(size)
        r.font.bold = b
        if color is not None:
            r.font.color.rgb = color
    return p


def md_segs(line):
    out = []
    for i, part in enumerate(re.split(r'\*\*', line)):
        if part:
            out.append((part.replace('`', ''), i % 2 == 1))
    return out


def add_image(d, path, caption):
    from PIL import Image
    w, h = Image.open(path).size
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run()
    if w >= h:
        run.add_picture(path, width=Cm(14))
    else:
        run.add_picture(path, height=Cm(16.5))
    add(d, caption, size=8.5, color=LGRAY, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)


def build_proposal(md, out, img_dir):
    d = new_doc()
    for ln in open(md, encoding='utf-8').read().split('\n'):
        s = ln.rstrip()
        if not s:
            continue
        m = re.fullmatch(r'\[\[IMG:([^|\]]+)\|([^\]]*)\]\]', s.strip())
        if m:
            add_image(d, os.path.join(img_dir, m.group(1)), m.group(2))
        elif s.startswith('# '):
            add(d, s[2:], size=16, bold=True, after=10)
        elif s.startswith('## '):
            add(d, s[3:], size=12.5, bold=True, before=12, after=4, keep_next=True)
        elif s.startswith('### '):
            add(d, s[4:], size=11, bold=True, after=3, page_break=True, keep_next=True)
        elif s.startswith('!! '):
            add(d, s[3:], size=10.5, bold=True, color=RED, after=2, fill='FDECEA')
        elif re.fullmatch(r'\*\*[^*]+\*\*', s) and len(s) < 20:
            add(d, s.strip('*'), size=10.5, bold=True, before=6, after=2, keep_next=True)
        elif s.startswith('- '):
            add(d, md_segs('· ' + s[2:]))
        else:
            add(d, md_segs(s))
    d.save(out)


def build_script(md, out):
    d = new_doc()
    first_ep = True
    for ln in open(md, encoding='utf-8').read().split('\n'):
        s = ln.rstrip()
        if not s:
            continue
        m = re.match(r'# (第\d+集)\s*$', s)
        if m:
            add(d, m.group(1), size=13, bold=True, page_break=True, after=6)
            first_ep = False
            continue
        if s.startswith('# '):
            add(d, s[2:], size=15, bold=True, after=6)
        elif s.startswith('## '):
            add(d, s[3:], size=11, bold=True, before=8, after=2, keep_next=True)
        elif s.startswith('>'):
            add(d, s.lstrip('> '), size=9, color=LGRAY, after=2)
        elif s.startswith('人物：'):
            add(d, s, size=9, color=LGRAY, after=4)
        elif s.startswith('△') or s.startswith('字幕：'):
            add(d, s, color=GRAY)
        elif s.startswith('- '):
            add(d, '· ' + s[2:])
        else:
            add(d, s)
    d.save(out)


if __name__ == '__main__':
    kind, md, out = sys.argv[1:4]
    if kind == 'proposal':
        build_proposal(md, out, sys.argv[4])
    else:
        build_script(md, out)
    print('saved', out)
