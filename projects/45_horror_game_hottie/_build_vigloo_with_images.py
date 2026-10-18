# -*- coding: utf-8 -*-
"""45번 Vigloo 기획안(한중 병기) 빌드 + 캐릭터 이미지 삽입.
  py -I -X utf8 _build_vigloo_with_images.py <spec.txt> <out.docx> <이미지 폴더>
tools/build_vigloo_proposal.py로 템플릿을 채운 뒤, 셀 안의 `[[IMG:파일|캡션]]` 줄을 그림 + 캡션으로 바꾼다.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'tools'))
from build_vigloo_proposal import build  # noqa: E402

IMG = re.compile(r'^\[\[IMG:([^|\]]+)\|([^\]]*)\]\]$')
FONT = 'Malgun Gothic'


def insert_images(path, img_dir):
    doc = Document(path)
    n = 0
    for row in doc.tables[0].rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                m = IMG.match(p.text.strip())
                if not m:
                    continue
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.keep_with_next = True
                f = os.path.join(img_dir, m.group(1))
                from PIL import Image
                w, h = Image.open(f).size
                run = p.add_run()
                if w >= h:
                    run.add_picture(f, width=Cm(10.0))
                else:
                    run.add_picture(f, height=Cm(11.0))
                cap = p.insert_paragraph_before('')  # 위치 잡기용 — 아래에서 순서 바꿈
                cap._p.getparent().remove(cap._p)
                p._p.addnext(cap._p)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(6)
                rr = cap.add_run(m.group(2))
                rr.font.name = FONT
                rr._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
                rr.font.size = Pt(8)
                rr.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
                n += 1
    doc.save(path)
    return n


if __name__ == '__main__':
    spec, out, img_dir = sys.argv[1:4]
    build(spec, out)
    print('images', insert_images(out, img_dir))
